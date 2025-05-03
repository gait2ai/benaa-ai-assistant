# embed_documents.py
#
# This script preprocesses documents in the data/ directory by:
# 1. Reading various document formats (TXT, DOCX, PDF, HTML, MD, XLSX, CSV, JSON)
# 2. Splitting content into meaningful chunks (200-300 words)
# 3. Creating vector embeddings using a multilingual SentenceTransformer model
# 4. Storing embeddings and metadata in a FAISS index for efficient retrieval

import os
import glob
import re
import json
import argparse
import logging
import time
import pandas as pd
from typing import List, Dict, Tuple, Any, Optional
from dataclasses import dataclass
from concurrent.futures import ThreadPoolExecutor

import numpy as np
import faiss
from tqdm import tqdm
from sentence_transformers import SentenceTransformer
from dotenv import load_dotenv

# Document processing libraries
import docx
import PyPDF2
import openpyxl
from bs4 import BeautifulSoup

# Configure logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s',
    handlers=[
        logging.FileHandler("embedding.log"),
        logging.StreamHandler()
    ]
)
logger = logging.getLogger(__name__)

# Load environment variables
load_dotenv()

# Constants
DEFAULT_DATA_FOLDER = "./data"
DEFAULT_OUTPUT_PATH = "./embeddings"
DEFAULT_CONFIG_FOLDER = "./config"
DEFAULT_SYSTEM_PROMPT_FILE = "system_prompt.txt"
DEFAULT_MIN_CHUNK_SIZE = 100  # Minimum number of characters for a chunk
DEFAULT_TARGET_CHUNK_SIZE = 250  # Target number of words per chunk
DEFAULT_MAX_CHUNK_SIZE = 350  # Maximum number of words per chunk
DEFAULT_PARAGRAPH_SEPARATOR = r'\n\s*\n'
DEFAULT_BATCH_SIZE = 32
DEFAULT_MODEL_NAME = "distiluse-base-multilingual-cased-v2"

@dataclass
class DocumentChunk:
    """Class to store document chunk information."""
    text: str
    source_file: str
    chunk_id: int
    is_system_prompt: bool = False  # Flag to identify system prompt chunks
    language: str = "en"  # Optional language identifier
    timestamp: str = ""  # Optional timestamp
    
    def to_dict(self) -> Dict[str, Any]:
        """Convert to dictionary for serialization."""
        return {
            "text": self.text,
            "source_file": self.source_file,
            "chunk_id": self.chunk_id,
            "is_system_prompt": self.is_system_prompt,
            "language": self.language,
            "timestamp": self.timestamp
        }

    @classmethod
    def from_dict(cls, data: Dict[str, Any]) -> 'DocumentChunk':
        """Create DocumentChunk from dictionary."""
        return cls(
            text=data["text"],
            source_file=data["source_file"],
            chunk_id=data["chunk_id"],
            is_system_prompt=data.get("is_system_prompt", False),
            language=data.get("language", "en"),
            timestamp=data.get("timestamp", "")
        )

class DocumentEmbedder:
    """Class to manage document embedding process."""
    
    def __init__(self, 
                 data_folder: str = DEFAULT_DATA_FOLDER, 
                 output_path: str = DEFAULT_OUTPUT_PATH,
                 config_folder: str = DEFAULT_CONFIG_FOLDER,
                 system_prompt_file: str = DEFAULT_SYSTEM_PROMPT_FILE,
                 model_name: str = DEFAULT_MODEL_NAME,
                 min_chunk_size: int = DEFAULT_MIN_CHUNK_SIZE,
                 target_chunk_size: int = DEFAULT_TARGET_CHUNK_SIZE,
                 max_chunk_size: int = DEFAULT_MAX_CHUNK_SIZE):
        """Initialize the document embedder.

        Args:
            data_folder: Path to the folder containing documents
            output_path: Path to save embeddings and metadata
            config_folder: Path to configuration folder
            system_prompt_file: Name of the system prompt file
            model_name: Name of the SentenceTransformer model
            min_chunk_size: Minimum character count for a valid text chunk
            target_chunk_size: Target word count for each chunk
            max_chunk_size: Maximum word count for each chunk
        """
        self.data_folder = data_folder
        self.output_path = output_path
        self.config_folder = config_folder
        self.system_prompt_file = system_prompt_file
        self.model_name = model_name
        self.min_chunk_size = min_chunk_size
        self.target_chunk_size = target_chunk_size
        self.max_chunk_size = max_chunk_size
        self.chunks = []
        self.embedding_model = None
        
        # Create output and config directories if they don't exist
        os.makedirs(output_path, exist_ok=True)
        os.makedirs(config_folder, exist_ok=True)
        
    def read_text_file(self, filepath: str) -> str:
        """Read content from a text file."""
        try:
            with open(filepath, 'r', encoding='utf-8') as file:
                return file.read()
        except UnicodeDecodeError:
            # Try with a different encoding if utf-8 fails
            try:
                with open(filepath, 'r', encoding='latin-1') as file:
                    return file.read()
            except Exception as e:
                logger.error(f"Error reading text file {filepath}: {e}")
                return ""
                
    def read_html_file(self, filepath: str) -> str:
        """Read and extract text from HTML file."""
        try:
            with open(filepath, 'r', encoding='utf-8') as file:
                html_content = file.read()
            soup = BeautifulSoup(html_content, 'html.parser')
            # Remove script and style elements
            for script in soup(["script", "style"]):
                script.extract()
            return soup.get_text(separator="\n")
        except Exception as e:
            logger.error(f"Error reading HTML file {filepath}: {e}")
            return ""

    def read_json_file(self, filepath: str) -> str:
        """Read and extract text from JSON file."""
        try:
            with open(filepath, 'r', encoding='utf-8') as file:
                data = json.load(file)
            
            # Handle common JSON structures
            if isinstance(data, dict):
                text_fields = []
                # Extract title if present
                if "title" in data:
                    text_fields.append(str(data["title"]))
                
                # Extract content fields if present
                for field in ["content", "text", "body", "description"]:
                    if field in data and data[field]:
                        if isinstance(data[field], str):
                            text_fields.append(data[field])
                        elif isinstance(data[field], list):
                            text_fields.extend([str(item) for item in data[field] if item])
                
                # If no specific fields found, convert the whole object to string
                if not text_fields:
                    return json.dumps(data, ensure_ascii=False, indent=2)
                
                return "\n\n".join(text_fields)
            
            elif isinstance(data, list):
                # Handle list of objects
                text_items = []
                for item in data:
                    if isinstance(item, dict):
                        item_text = []
                        for k, v in item.items():
                            if isinstance(v, str) and len(v) > 10:  # Only include substantial text
                                item_text.append(f"{k}: {v}")
                        if item_text:
                            text_items.append("\n".join(item_text))
                    elif isinstance(item, str):
                        text_items.append(item)
                
                return "\n\n".join(text_items)
            
            return str(data)
            
        except Exception as e:
            logger.error(f"Error reading JSON file {filepath}: {e}")
            return ""

    def read_docx_file(self, filepath: str) -> str:
        """Read and extract text from DOCX file."""
        try:
            doc = docx.Document(filepath)
            return "\n".join([para.text for para in doc.paragraphs if para.text])
        except Exception as e:
            logger.error(f"Error reading DOCX file {filepath}: {e}")
            return ""

    def read_pdf_file(self, filepath: str) -> str:
        """Read and extract text from PDF file."""
        try:
            text = []
            with open(filepath, 'rb') as file:
                reader = PyPDF2.PdfReader(file)
                for page in reader.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text.append(page_text)
            return "\n\n".join(text)
        except Exception as e:
            logger.error(f"Error reading PDF file {filepath}: {e}")
            return ""

    def read_xlsx_file(self, filepath: str) -> str:
        """Read and extract text from XLSX file."""
        try:
            wb = openpyxl.load_workbook(filepath, data_only=True)
            text = []
            
            for sheet in wb.worksheets:
                sheet_text = []
                # Get sheet name as a heading
                sheet_text.append(f"# {sheet.title}")
                
                for row in sheet.iter_rows(values_only=True):
                    # Skip empty rows
                    if not any(cell for cell in row):
                        continue
                    # Convert row to text
                    row_text = " | ".join([str(cell) if cell is not None else "" for cell in row])
                    sheet_text.append(row_text)
                
                text.append("\n".join(sheet_text))
            
            return "\n\n".join(text)
        except Exception as e:
            logger.error(f"Error reading XLSX file {filepath}: {e}")
            return ""

    def read_csv_file(self, filepath: str) -> str:
        """Read and extract text from CSV file."""
        try:
            # Try reading with pandas - handles various CSV formats
            df = pd.read_csv(filepath, encoding='utf-8', on_bad_lines='skip')
            
            # Get column names as a header row
            headers = " | ".join(df.columns)
            
            # Convert dataframe rows to text
            rows = []
            for _, row in df.iterrows():
                row_text = " | ".join([str(val) if not pd.isna(val) else "" for val in row])
                rows.append(row_text)
            
            # Combine header and rows
            text = f"{headers}\n" + "\n".join(rows)
            return text
            
        except UnicodeDecodeError:
            # Try with a different encoding if utf-8 fails
            try:
                df = pd.read_csv(filepath, encoding='latin-1', on_bad_lines='skip')
                headers = " | ".join(df.columns)
                rows = []
                for _, row in df.iterrows():
                    row_text = " | ".join([str(val) if not pd.isna(val) else "" for val in row])
                    rows.append(row_text)
                text = f"{headers}\n" + "\n".join(rows)
                return text
            except Exception as e:
                logger.error(f"Error reading CSV file with alternative encoding {filepath}: {e}")
                return ""
        except Exception as e:
            logger.error(f"Error reading CSV file {filepath}: {e}")
            return ""

    def read_file(self, filepath: str) -> str:
        """Read file content based on file extension."""
        try:
            ext = os.path.splitext(filepath)[1].lower()
            
            if ext in ['.txt', '.md']:
                return self.read_text_file(filepath)
            elif ext == '.html':
                return self.read_html_file(filepath)
            elif ext == '.json':
                return self.read_json_file(filepath)
            elif ext == '.docx':
                return self.read_docx_file(filepath)
            elif ext == '.pdf':
                return self.read_pdf_file(filepath)
            elif ext == '.xlsx':
                return self.read_xlsx_file(filepath)
            elif ext == '.csv':
                return self.read_csv_file(filepath)
            else:
                logger.warning(f"Unsupported file type: {filepath}")
                return ""
        except Exception as e:
            logger.error(f"Error reading file {filepath}: {e}")
            return ""

    def clean_text(self, text: str) -> str:
        """Clean and normalize text."""
        # Replace multiple whitespace with single space
        text = re.sub(r'\s+', ' ', text)
        # Remove control characters
        text = re.sub(r'[\x00-\x1F\x7F]', '', text)
        # Trim whitespace
        return text.strip()
    
    def split_into_chunks(self, text: str, source_file: str, is_system_prompt: bool = False) -> List[DocumentChunk]:
        """Split text into chunks of appropriate size.
        
        This method splits text into chunks of around 200-300 words, preferring to split
        at paragraph or sentence boundaries.
        """
        if not text:
            return []
        
        # For system prompts, we keep them as is
        if is_system_prompt:
            clean_text = self.clean_text(text)
            if not clean_text:
                return []
            
            # Create a single chunk with the entire system prompt
            return [DocumentChunk(
                text=clean_text,
                source_file=source_file,
                chunk_id=0,
                is_system_prompt=True,
                timestamp=time.strftime("%Y-%m-%d %H:%M:%S")
            )]
        
        # For regular documents, split into appropriate chunks
        # First, split by paragraphs
        paragraphs = re.split(DEFAULT_PARAGRAPH_SEPARATOR, text)
        paragraphs = [self.clean_text(p) for p in paragraphs if self.clean_text(p)]
        
        # Remove any empty paragraphs after cleaning
        paragraphs = [p for p in paragraphs if p]
        
        if not paragraphs:
            return []
            
        chunks = []
        current_chunk = []
        current_word_count = 0
        chunk_id = 0
        
        for para in paragraphs:
            para_words = para.split()
            para_word_count = len(para_words)
            
            # If paragraph is very large, split it into sentences
            if para_word_count > self.max_chunk_size:
                sentences = re.split(r'(?<=[.!?])\s+', para)
                for sentence in sentences:
                    if not sentence.strip():
                        continue
                        
                    sentence_words = sentence.split()
                    sentence_word_count = len(sentence_words)
                    
                    # If adding this sentence would exceed max chunk size, start a new chunk
                    if current_word_count + sentence_word_count > self.max_chunk_size and current_chunk:
                        chunk_text = " ".join(current_chunk)
                        if len(chunk_text) >= self.min_chunk_size:
                            chunks.append(DocumentChunk(
                                text=chunk_text,
                                source_file=source_file,
                                chunk_id=chunk_id,
                                timestamp=time.strftime("%Y-%m-%d %H:%M:%S")
                            ))
                            chunk_id += 1
                        current_chunk = []
                        current_word_count = 0
                    
                    current_chunk.append(sentence)
                    current_word_count += sentence_word_count
                    
                    # If we've reached the target size after adding a sentence, start a new chunk
                    if current_word_count >= self.target_chunk_size:
                        chunk_text = " ".join(current_chunk)
                        if len(chunk_text) >= self.min_chunk_size:
                            chunks.append(DocumentChunk(
                                text=chunk_text,
                                source_file=source_file,
                                chunk_id=chunk_id,
                                timestamp=time.strftime("%Y-%m-%d %H:%M:%S")
                            ))
                            chunk_id += 1
                        current_chunk = []
                        current_word_count = 0
            else:
                # If adding this paragraph would exceed max chunk size, start a new chunk
                if current_word_count + para_word_count > self.max_chunk_size and current_chunk:
                    chunk_text = " ".join(current_chunk)
                    if len(chunk_text) >= self.min_chunk_size:
                        chunks.append(DocumentChunk(
                            text=chunk_text,
                            source_file=source_file,
                            chunk_id=chunk_id,
                            timestamp=time.strftime("%Y-%m-%d %H:%M:%S")
                        ))
                        chunk_id += 1
                    current_chunk = []
                    current_word_count = 0
                
                current_chunk.append(para)
                current_word_count += para_word_count
                
                # If we've reached the target size after adding a paragraph, start a new chunk
                if current_word_count >= self.target_chunk_size:
                    chunk_text = " ".join(current_chunk)
                    if len(chunk_text) >= self.min_chunk_size:
                        chunks.append(DocumentChunk(
                            text=chunk_text,
                            source_file=source_file,
                            chunk_id=chunk_id,
                            timestamp=time.strftime("%Y-%m-%d %H:%M:%S")
                        ))
                        chunk_id += 1
                    current_chunk = []
                    current_word_count = 0
        
        # Add any remaining content as a final chunk
        if current_chunk:
            chunk_text = " ".join(current_chunk)
            if len(chunk_text) >= self.min_chunk_size:
                chunks.append(DocumentChunk(
                    text=chunk_text,
                    source_file=source_file,
                    chunk_id=chunk_id,
                    timestamp=time.strftime("%Y-%m-%d %H:%M:%S")
                ))
        
        return chunks

    def read_system_prompt(self) -> Optional[str]:
        """Read the system prompt file if it exists."""
        system_prompt_path = os.path.join(self.config_folder, self.system_prompt_file)
        
        if os.path.exists(system_prompt_path):
            logger.info(f"Reading system prompt from {system_prompt_path}")
            try:
                content = self.read_text_file(system_prompt_path)
                if content:
                    logger.info(f"System prompt loaded successfully ({len(content)} characters)")
                    return content
                else:
                    logger.warning(f"System prompt file exists but is empty: {system_prompt_path}")
            except Exception as e:
                logger.error(f"Error reading system prompt file: {e}")
        else:
            logger.info(f"No system prompt file found at {system_prompt_path}")
            
        return None

    def load_documents(self) -> None:
        """Load all documents from data folder and prepare chunks."""
        logger.info(f"Loading documents from {self.data_folder}")
        
        # First, check for system prompt and add it as a special chunk if exists
        system_prompt = self.read_system_prompt()
        if system_prompt:
            # Create a special system prompt chunk with a distinctive source name
            system_chunks = self.split_into_chunks(
                system_prompt, 
                source_file="__SYSTEM_PROMPT__",
                is_system_prompt=True
            )
            
            self.chunks.extend(system_chunks)
            logger.info(f"Added {len(system_chunks)} system prompt chunks")
        
        # Get all supported files
        file_patterns = [
            "**/*.txt", "**/*.md", "**/*.html", 
            "**/*.json", "**/*.docx", "**/*.pdf", 
            "**/*.xlsx", "**/*.csv"
        ]
        
        all_files = []
        for pattern in file_patterns:
            glob_pattern = os.path.join(self.data_folder, pattern)
            all_files.extend(glob.glob(glob_pattern, recursive=True))
        
        logger.info(f"Found {len(all_files)} files")
        if not all_files:
            logger.warning(f"No documents found in {self.data_folder}")
            if not system_prompt:
                # If there's no system prompt either, we have nothing to embed
                return
        
        # Process files in parallel
        chunks = []
        with ThreadPoolExecutor() as executor:
            futures = []
            
            for filepath in all_files:
                logger.debug(f"Processing file: {filepath}")
                relative_path = os.path.relpath(filepath, self.data_folder)
                
                # Submit file processing task
                future = executor.submit(self.process_file, filepath, relative_path)
                futures.append(future)
            
            # Collect results with a progress bar
            for future in tqdm(futures, desc="Processing documents"):
                file_chunks = future.result()
                chunks.extend(file_chunks)
        
        self.chunks.extend(chunks)
        logger.info(f"Created {len(self.chunks)} document chunks total (including system prompt if present)")

    def process_file(self, filepath: str, relative_path: str) -> List[DocumentChunk]:
        """Process a single file and return chunks."""
        content = self.read_file(filepath)
        if not content:
            logger.warning(f"No content extracted from file: {filepath}")
            return []
            
        logger.debug(f"Extracted {len(content)} characters from {filepath}")
        return self.split_into_chunks(content, relative_path)

    def load_embedding_model(self) -> None:
        """Load the SentenceTransformer model."""
        logger.info(f"Loading embedding model: {self.model_name}")
        start_time = time.time()
        self.embedding_model = SentenceTransformer(self.model_name)
        logger.info(f"Model loaded in {time.time() - start_time:.2f} seconds")

    def create_embeddings(self) -> Tuple[np.ndarray, List[DocumentChunk]]:
        """Create embeddings for all document chunks."""
        if not self.chunks:
            logger.warning("No document chunks to embed")
            return np.array([]), []
        
        if not self.embedding_model:
            self.load_embedding_model()
        
        texts = [chunk.text for chunk in self.chunks]
        
        logger.info(f"Creating embeddings for {len(texts)} chunks")
        start_time = time.time()
        
        # Create embeddings in batches
        embeddings = []
        batch_size = DEFAULT_BATCH_SIZE
        
        for i in range(0, len(texts), batch_size):
            batch_texts = texts[i:i+batch_size]
            batch_embeddings = self.embedding_model.encode(batch_texts, show_progress_bar=False)
            embeddings.append(batch_embeddings)
            
        embeddings_array = np.vstack(embeddings)
        
        logger.info(f"Created embeddings with shape {embeddings_array.shape} in {time.time() - start_time:.2f} seconds")
        return embeddings_array, self.chunks

    def build_faiss_index(self, embeddings: np.ndarray) -> faiss.Index:
        """Build a FAISS index from embeddings."""
        vector_dimension = embeddings.shape[1]
        
        # Create a simple L2 distance index
        index = faiss.IndexFlatL2(vector_dimension)
        
        # Add vectors to the index
        faiss.normalize_L2(embeddings)  # Normalize for cosine similarity
        index.add(embeddings)
        
        logger.info(f"Built FAISS index with {index.ntotal} vectors")
        return index

    def save_faiss_index(self, index: faiss.Index, filename: str = "document_index.faiss") -> None:
        """Save FAISS index to disk."""
        index_path = os.path.join(self.output_path, filename)
        faiss.write_index(index, index_path)
        logger.info(f"Saved FAISS index to {index_path}")

    def save_metadata(self, chunks: List[DocumentChunk], filename: str = "document_metadata.json") -> None:
        """Save document chunks metadata to disk."""
        metadata_path = os.path.join(self.output_path, filename)
        
        # Convert chunks to dictionaries for serialization
        chunks_data = [chunk.to_dict() for chunk in chunks]
        
        with open(metadata_path, 'w', encoding='utf-8') as f:
            json.dump(chunks_data, f, ensure_ascii=False, indent=2)
            
        logger.info(f"Saved metadata for {len(chunks)} chunks to {metadata_path}")

    def save_model_info(self, filename: str = "model_info.json") -> None:
        """Save embedding model information."""
        model_info_path = os.path.join(self.output_path, filename)
        
        model_info = {
            "model_name": self.model_name,
            "vector_dimension": self.embedding_model.get_sentence_embedding_dimension(),
            "created_at": time.strftime("%Y-%m-%d %H:%M:%S"),
            "has_system_prompt": any(chunk.is_system_prompt for chunk in self.chunks),
            "target_chunk_size": self.target_chunk_size,
            "max_chunk_size": self.max_chunk_size,
            "min_chunk_size": self.min_chunk_size
        }
        
        with open(model_info_path, 'w', encoding='utf-8') as f:
            json.dump(model_info, f, indent=2)
            
        logger.info(f"Saved model information to {model_info_path}")

    def process(self) -> None:
        """Run the entire document embedding process."""
        try:
            logger.info("Starting document embedding process")
            
            # Load and process documents
            self.load_documents()
            if not self.chunks:
                logger.error("No valid document chunks found. Process aborted.")
                return
            
            # Create embeddings
            embeddings, chunks = self.create_embeddings()
            if len(embeddings) == 0:
                logger.error("Failed to create embeddings. Process aborted.")
                return
            
            # Build and save FAISS index
            index = self.build_faiss_index(embeddings)
            self.save_faiss_index(index)
            
            # Save metadata and model info
            self.save_metadata(chunks)
            self.save_model_info()
            
            logger.info("Document embedding process completed successfully")
            
        except Exception as e:
            logger.error(f"Error during document embedding process: {e}", exc_info=True)


def main():
    """Main function to run the document embedding process."""
    parser = argparse.ArgumentParser(description="Embed documents for RAG retrieval")
    parser.add_argument("--data", type=str, default=DEFAULT_DATA_FOLDER,
                        help=f"Path to the documents folder (default: {DEFAULT_DATA_FOLDER})")
    parser.add_argument("--output", type=str, default=DEFAULT_OUTPUT_PATH,
                        help=f"Path to save embeddings and metadata (default: {DEFAULT_OUTPUT_PATH})")
    parser.add_argument("--config", type=str, default=DEFAULT_CONFIG_FOLDER,
                        help=f"Path to configuration folder (default: {DEFAULT_CONFIG_FOLDER})")
    parser.add_argument("--prompt-file", type=str, default=DEFAULT_SYSTEM_PROMPT_FILE,
                        help=f"Name of system prompt file (default: {DEFAULT_SYSTEM_PROMPT_FILE})")
    parser.add_argument("--model", type=str, default=DEFAULT_MODEL_NAME,
                        help=f"SentenceTransformer model name (default: {DEFAULT_MODEL_NAME})")
    parser.add_argument("--min-chunk-size", type=int, default=DEFAULT_MIN_CHUNK_SIZE,
                        help=f"Minimum character length for chunks (default: {DEFAULT_MIN_CHUNK_SIZE})")
    parser.add_argument("--target-chunk-size", type=int, default=DEFAULT_TARGET_CHUNK_SIZE,
                        help=f"Target word count for chunks (default: {DEFAULT_TARGET_CHUNK_SIZE})")
    parser.add_argument("--max-chunk-size", type=int, default=DEFAULT_MAX_CHUNK_SIZE,
                        help=f"Maximum word count for chunks (default: {DEFAULT_MAX_CHUNK_SIZE})")
    parser.add_argument("--log-level", type=str, default="INFO",
                        choices=["DEBUG", "INFO", "WARNING", "ERROR"],
                        help="Logging level (default: INFO)")
    
    args = parser.parse_args()
    
    # Set log level
    logging.getLogger().setLevel(getattr(logging, args.log_level))
    
    # Create and run the document embedder
    embedder = DocumentEmbedder(
        data_folder=args.data,
        output_path=args.output,
        config_folder=args.config,
        system_prompt_file=args.prompt_file,
        model_name=args.model,
        min_chunk_size=args.min_chunk_size,
        target_chunk_size=args.target_chunk_size,
        max_chunk_size=args.max_chunk_size
    )
    
    embedder.process()


if __name__ == "__main__":
    main()